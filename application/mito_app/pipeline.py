"""End-to-end pipeline: nucleus-seeded cell ROIs -> per-cell MiNA on the red channel.

Inputs are three images of the same field (ImageJ RGB TIFF exports, signal in the
matching channel, or single-channel images):
  red   = mitochondria (MiNA is run on it)
  green = protein of interest (mean intensity per cell; its weak autofluorescence
          also outlines the cells for the ROI step)
  blue  = nuclei (one in-focus nucleus = one cell)
"""
import csv, os, re
from dataclasses import asdict, dataclass, fields, replace

import numpy as np

from . import cell_roi, imgio, nuclei, green_cells, mina, mito_objects, stats
from . import green as green_q


@dataclass
class Params:
    min_area_um2: float = 250.0       # smallest cell ROI kept (cell_roi --min-area-um2)
    binuc_tau: float = 0.0            # border contrast below which neighbours are "weak border" pairs
    exclude_binucleate: bool = False  # drop weak-border pairs instead of keeping them split per nucleus
    include_edge_cells: bool = False  # also run MiNA on cells touching the frame
    pixel_size_um: float = 0.0        # 0 = read from the TIFF (falls back to the reference size)
    puncta_sensitivity: float = 1.0   # multiplies the green puncta threshold (>1 = fewer, brighter puncta)
    min_mito_area_um2: float = 0.05   # smallest mito object in the per-mito table
    mito_method: str = 'split'        # 'split' = adaptive threshold + watershed objects, 'otsu' = MiNA classic
    green_pos_percent: float = 2.0    # cell is green-positive if bright green covers >= this % of its cytoplasm
    dim_nuclei: bool = True           # out-of-focus nuclei also seed (and own) a cell
    trim_edge_cells: bool = True      # cut a frame-touching tip off along a mito-free line so the cell is kept
    nuclei_method: str = 'texture_merge'  # nucleus detection: 'texture_merge', 'texture' or 'otsu' (mito_app/nuclei.py)
    # manual thresholds (0 = automatic, computed per image); see THRESHOLDS and mito_app/thresholds.py
    thr_nuclei: float = 0.0           # nucleus candidates: flattened (texture) or smoothed (otsu) blue
    cell_fg_level: float = 0.12       # cell area: mito density (red at 3 um / its 99th percentile) above this
    thr_mito: float = 0.0             # mito mask: preprocessed red (rolling ball 1.5 um + sigma 0.7 px)
    thr_green_bright: float = 0.0     # green+ cells: bright green, smoothed auto-levelled green
    thr_puncta: float = 0.0           # green puncta: top-hat of the background-subtracted green


# (Params field, label, unit, image it applies to, what "auto" means)
THRESHOLDS = (
    ('thr_nuclei', 'Nuclei', 'blue', 'flattened blue (texture) / smoothed blue (Otsu)', '0.5 × Otsu (texture), Otsu'),
    ('cell_fg_level', 'Cell area', '× p99', 'mito density (red at 3 µm ÷ its 99th percentile)', 'fixed 0.12'),
    ('thr_mito', 'Mitochondria', 'red', 'preprocessed red (rolling ball 1.5 µm, σ 0.7 px)',
     'per cell: local mean AND 0.5 × Otsu (split) / Otsu of raw red (MiNA classic)'),
    ('thr_green_bright', 'Green+ (bright green)', 'green AL', 'smoothed auto-levelled green',
     'two Otsu steps, ≥ 3 × cytoplasm level'),
    ('thr_puncta', 'Green puncta', 'top-hat', 'top-hat of background-subtracted green',
     'max(Otsu, median + 6 MAD) × sensitivity'),
)
PARAM_FIELDS = {f.name for f in fields(Params)}


def sample_name(red_path):
    """'Processed_New01_Ch1_Red.tif' -> 'Processed_New01'; 'X_Ch1_Red_Cropped_ROI1.tif' -> 'X_Cropped_ROI1';
    a CZI channel 'New-01.czi::ch0' -> 'New-01'."""
    czi, ch = imgio.split_ref(red_path)
    if ch is not None:
        return os.path.splitext(os.path.basename(czi))[0]
    stem = os.path.splitext(os.path.basename(red_path))[0]
    out = re.sub(r'[_\- ]*(ch\d+)?[_\- ]*(red|mito)$', '', stem, flags=re.I)
    if out == stem:  # channel tag in the middle of the name
        out = re.sub(r'(^|[_\- ]+)(ch\d+[_\- ]*)?red(?=[_\- ])', '', stem, count=1, flags=re.I)
    return out.strip('_- ') or stem


CHANNEL_RX = re.compile(r'(ch\d+[_\- ]*)?(red|green|blue)', re.I)


def find_siblings(path):
    """Guess the other two channel files from one file name (…Ch1_Red / …Ch2_Green / …Ch3_Blue)."""
    d, name = os.path.split(path)
    m = CHANNEL_RX.search(name)
    if not m:
        return {}
    files = os.listdir(d or '.')
    out = {}
    for color in ('red', 'green', 'blue'):
        rx = re.compile(re.escape(name[:m.start()]) + r'(ch\d+[_\- ]*)?' + color + re.escape(name[m.end():]), re.I)
        hits = [f for f in files if rx.fullmatch(f)]
        if len(hits) == 1:
            out[color] = os.path.join(d, hits[0])
    return out


def find_sets(paths, czi_roles=None):
    """Group TIFF files (or the TIFFs inside folders, searched recursively) into red/green/blue image sets
    by file name. A CZI file is one set by itself: its channels are given the roles in `czi_roles`
    ({'red': mito channel, 'green': protein channel, 'blue': nucleus channel}, 0-based indices) or, when that is
    None, the roles guessed from its channel names (imgio.guess_roles); the set then carries 'czi' and 'roles'.

    Returns (sets, leftover): sets is a list of {'red', 'green', 'blue', 'dataset', 'name', 'root'} dicts in name
    order, where dataset is the name of the folder that was added (or the file's own folder for single files) and
    root that folder's path;
    leftover the channel files (names with red/green/blue) that could not be placed in a complete set."""
    files = {}  # path -> (dataset, folder the dataset name comes from)
    for p in paths:
        p = os.path.abspath(p)
        if os.path.isdir(p):
            ds = os.path.basename(p.rstrip(os.sep))
            for d, dirs, fs in os.walk(p):
                dirs[:] = sorted(x for x in dirs if not x.startswith('.'))
                for f in sorted(fs):
                    if f.lower().endswith(('.tif', '.tiff', imgio.CZI_EXT)) and not f.startswith('.'):
                        files.setdefault(os.path.join(d, f), (ds, p))
        elif os.path.isfile(p):
            files.setdefault(p, (os.path.basename(os.path.dirname(p)), os.path.dirname(p)))
    sets, used, bad_czi = {}, set(), []
    for f, (ds, root) in files.items():
        if imgio.is_czi(f):
            try:
                n = len(imgio.czi_info(f)['channels'])
                roles = czi_roles or imgio.guess_roles(imgio.czi_info(f)['channels'])
                ok = len(roles) == 3 and max(roles.values()) < n
            except Exception:
                ok = False
            if ok:
                used.add(f)
                s = imgio.czi_set(f, roles, ds, root)
                s['_root'] = s.pop('root')
                sets[s['red']] = s
            else:
                bad_czi.append(f)
            continue
        s = find_siblings(f)
        if len(s) == 3:
            used.update(s.values())
            sets[s['red']] = dict(s, dataset=ds, name=sample_name(s['red']), _root=root)
    # the same file name in several subfolders of one dataset: prefix the subfolder so output folders differ
    seen = {}
    for s in sets.values():
        seen.setdefault((s['dataset'], s['name']), []).append(s)
    for group in seen.values():
        if len(group) > 1:
            for s in group:
                sub = os.path.relpath(os.path.dirname(s.get('czi') or s['red']), s['_root'])
                if sub != '.':
                    s['name'] = f"{sub.replace(os.sep, '_')}_{s['name']}"
    for s in sets.values():
        s['root'] = s.pop('_root')
    # Merged / labels / binary TIFFs etc. are not channel files and are not reported
    leftover = [f for f in files if f not in used and CHANNEL_RX.search(os.path.basename(f))] + bad_czi
    return [sets[k] for k in sorted(sets)], leftover


def safe_name(s):
    """Folder-safe version of a dataset / preset / sample name."""
    return re.sub(r'[\\/:*?"<>|]+', '_', str(s)).strip(' .') or 'unnamed'


def job_outdir(base, dataset, preset, name):
    """<output>/<dataset>/<preset>-<dataset>/<sample>/"""
    d, p = safe_name(dataset), safe_name(preset)
    return os.path.join(base, d, f'{p}-{d}', safe_name(name))


def default_outdir(roots):
    """<common parent folder of the input folders>/dataset, or '' if they share no folder.
    One input folder (or nested ones): the folder above it."""
    roots = sorted({os.path.abspath(r) for r in roots if r})
    if not roots:
        return ''
    try:
        common = os.path.commonpath(roots)
    except ValueError:  # different drives
        return ''
    if common in roots:
        common = os.path.dirname(common)
    return os.path.join(common, 'dataset') if common and common != os.path.dirname(common) else ''


GROUPS_FILE = 'groups.csv'


def _read_group_file(path):
    with open(path, newline='') as f:
        return [r for r in csv.DictReader(f) if r.get('folder')]


def write_groups(base, entries):
    """Record the group of each sample folder in <base>/groups.csv (merged with what is already there).
    `entries`: dicts with outdir (the sample folder), dataset, sample and group."""
    path = os.path.join(base, GROUPS_FILE)
    rows = {r['folder']: r for r in _read_group_file(path)} if os.path.exists(path) else {}
    for e in entries:
        rel = os.path.relpath(e['outdir'], base).replace(os.sep, '/')
        rows[rel] = dict(folder=rel, dataset=e.get('dataset', ''), sample=e.get('sample', e.get('name', '')),
                         group=str(e.get('group', '')).strip())
    os.makedirs(base, exist_ok=True)
    write_csv(path, sorted(rows.values(), key=lambda r: r['folder']), ('folder', 'dataset', 'sample', 'group'))


def read_groups(root):
    """{absolute sample folder: group} from every groups.csv under `root`; a file nearer the root wins."""
    files = []
    for d, dirs, fs in os.walk(root):
        dirs.sort()
        if GROUPS_FILE in fs:
            files.append(d)
    out = {}
    for d in sorted(files, key=lambda x: -x.count(os.sep)):
        for r in _read_group_file(os.path.join(d, GROUPS_FILE)):
            if r.get('group', '').strip():
                out[os.path.normpath(os.path.join(d, r['folder']))] = r['group'].strip()
    return out


def write_group_tables(results):
    """Pool the per-cell / per-mito tables of all samples in each <preset>-<dataset> folder
    (`<preset>-<dataset>_all_cells.csv`, `…_all_mito.csv`, `…_all_results.xlsx`)."""
    groups = {}
    for r in results:
        groups.setdefault(os.path.dirname(r['outdir']), []).append(r)
    for d, rs in groups.items():
        tag = os.path.basename(d)
        cells = [dict(sample=r['name'], **row) for r in rs for row in r['rows']]
        mito = [dict(sample=r['name'], **row) for r in rs for row in r['mito_rows']]
        write_csv(os.path.join(d, f'{tag}_all_cells.csv'), cells)
        write_csv(os.path.join(d, f'{tag}_all_mito.csv'), mito)
        sheets = []
        for g_key, g_tag in GREEN_GROUPS:
            gc = [r for r in cells if r.get('green_status') == g_key]
            gm = [r for r in mito if r.get('green_status') == g_key]
            write_csv(os.path.join(d, f'{tag}_all_cells_{g_tag}.csv'), gc, cells[0] if cells else None)
            write_csv(os.path.join(d, f'{tag}_all_mito_{g_tag}.csv'), gm, mito[0] if mito else None)
            sheets += [(f'cells_{g_tag}', gc), (f'mito_{g_tag}', gm)]
        write_xlsx(os.path.join(d, f'{tag}_all_results.xlsx'), sheets + [('all_cells', cells), ('all_mito', mito)])


GROUP_EXPORT_DIR = 'groups'


def write_group_exports(root, out=None, control=None, log=print):
    """Group-wise data export of every sample under `root` (groups from groups.csv, default = dataset):

      <out>/<group>/<group>_cells.csv, _mito.csv            all cells / mito objects of the group
      <out>/<group>/<group>_cells_green_pos.csv, …_neg.csv   the same per green status (and for mito)
      <out>/<group>/<group>_per_image.csv                   one row per image: mean of each cell metric
      <out>/<group>/<group>_correlations.csv                regression / Spearman of every metric pair
      <out>/<group>/<group>_results.xlsx                    all of the above as sheets
      <out>/group_comparison.xlsx (+ _cells.csv, _cells_per_image.csv)
          every metric x group: n, mean, SD, SEM, 95 % CI, median, IQR, the across-group tests and, with a
          `control` group, each group vs control (Holm-adjusted); for all / green+ / green- cells, per cell,
          per image and per mito object
      <out>/groups_overview.csv                             images, cells, green+ cells, mito objects per group
    `out` defaults to <root>/groups. Returns the output folder, or '' when there is nothing to export."""
    cells, mito, samples = load_results(root)
    if not samples:
        return ''
    out = out or os.path.join(root, GROUP_EXPORT_DIR)
    os.makedirs(out, exist_ok=True)
    groups = sorted({s['group'] for s in samples})
    xs_c = stats.available(stats.CELL_X + stats.CELL_Y, cells)
    xs_m = stats.available(stats.MITO_X + stats.MITO_Y, mito)
    keys_c = [k for k, _ in xs_c]
    overview = []
    for g in groups:
        gd = os.path.join(out, safe_name(g)); os.makedirs(gd, exist_ok=True)
        tag = safe_name(g)
        gc = [r for r in cells if r['group'] == g]; gm = [r for r in mito if r['group'] == g]
        img = stats.sample_means(gc, keys_c)
        corr = stats.correlation_table(gc, gm)
        sheets = [('cells', gc), ('mito', gm), ('per_image', img), ('correlations', corr)]
        write_csv(os.path.join(gd, f'{tag}_cells.csv'), gc)
        write_csv(os.path.join(gd, f'{tag}_mito.csv'), gm)
        write_csv(os.path.join(gd, f'{tag}_per_image.csv'), img)
        write_csv(os.path.join(gd, f'{tag}_correlations.csv'), corr)
        for g_key, g_tag in GREEN_GROUPS:
            sc = [r for r in gc if r.get('green_status') == g_key]; sm = [r for r in gm if r.get('green_status') == g_key]
            write_csv(os.path.join(gd, f'{tag}_cells_{g_tag}.csv'), sc, gc[0] if gc else None)
            write_csv(os.path.join(gd, f'{tag}_mito_{g_tag}.csv'), sm, gm[0] if gm else None)
            sheets += [(f'cells_{g_tag}', sc), (f'mito_{g_tag}', sm)]
        write_xlsx(os.path.join(gd, f'{tag}_results.xlsx'), sheets)
        overview.append(dict(group=g, images=len({(s['dataset'], s['preset'], s['sample']) for s in samples
                                                   if s['group'] == g}),
                             cells=len(gc), green_pos_cells=sum(r.get('green_status') == 'positive' for r in gc),
                             mito_objects=len(gm)))
    write_csv(os.path.join(out, 'groups_overview.csv'), overview)
    comp = []
    for label, flt in (('all', ''), ('green_pos', 'positive'), ('green_neg', 'negative')):
        c = [r for r in cells if not flt or r.get('green_status') == flt]
        m = [r for r in mito if not flt or r.get('green_status') == flt]
        per_cell = stats.group_summary_long(c, xs_c, groups, control)
        per_img = stats.group_summary_long(stats.sample_means(c, keys_c), xs_c, groups, control)
        per_mito = stats.group_summary_long(m, xs_m, groups, control)
        comp += [(f'cells_{label}', per_cell), (f'images_{label}', per_img), (f'mito_{label}', per_mito)]
        if label == 'all':
            write_csv(os.path.join(out, 'group_comparison_cells.csv'), per_cell)
            write_csv(os.path.join(out, 'group_comparison_cells_per_image.csv'), per_img)
    write_xlsx(os.path.join(out, 'group_comparison.xlsx'), [('overview', overview)] + comp)
    log(f'Group export: {len(groups)} group(s) -> {out}')
    return out


def _num(v):
    try:
        return float(v) if v not in ('', 'True', 'False') else v
    except ValueError:
        return v


def read_csv(path):
    with open(path, newline='') as f:
        return [{k: _num(v) for k, v in r.items()} for r in csv.DictReader(f)]


def load_results(root):
    """Every sample under `root` (any depth): rows of <sample>_per_cell.csv and <sample>_per_mito.csv with
    dataset / preset / sample columns taken from the <dataset>/<preset>-<dataset>/<sample>/ layout, and a group
    column from groups.csv (default: the dataset).
    Returns (cell_rows, mito_rows, samples): samples = one dict per sample (folder + those columns)."""
    cells, mito, samples = [], [], []
    groups = read_groups(root)
    for d, dirs, files in os.walk(root):
        dirs.sort()
        for f in sorted(files):
            if not f.endswith('_per_cell.csv') or f.endswith('_mina_per_cell.csv'):
                continue
            name = f[:-len('_per_cell.csv')]
            up = os.path.basename(os.path.dirname(d))
            dataset = os.path.basename(os.path.dirname(os.path.dirname(d)))
            preset = up[:-len(dataset) - 1] if dataset and up.endswith('-' + dataset) else up
            meta = dict(group=groups.get(os.path.normpath(d), dataset), dataset=dataset, preset=preset, sample=name)
            cells += [dict(meta, **r) for r in read_csv(os.path.join(d, f))]
            mp = os.path.join(d, f'{name}_per_mito.csv')
            if os.path.exists(mp):
                mito += [dict(meta, **r) for r in read_csv(mp)]
            samples.append(dict(folder=d, **meta))
    return cells, mito, samples


def _as_rgb(a, g):
    """cell_roi overlays expect an RGB uint8 array; wrap single-channel input as green."""
    if a.ndim == 3 and a.shape[-1] == 3:
        return a
    g8 = np.clip(g / max(np.percentile(g, 99.5), 1) * 255, 0, 255).astype(np.uint8)
    return np.dstack([np.zeros_like(g8), g8, np.zeros_like(g8)])


def _cell_stage(red_path, green_path, blue_path, p, log):
    """Steps 1-2: load the images, find the nuclei and the cell ROIs (before any manual edit)."""
    log('1/5  Loading images')
    a, g, px = cell_roi.load_green(green_path)
    px_source = 'CZI' if imgio.is_czi(imgio.split_ref(green_path)[0]) else 'TIFF'
    if p.pixel_size_um > 0:
        px, px_source = p.pixel_size_um, 'entered'
    elif not px:  # no calibration in the file
        px, px_source = cell_roi.REF_PX_UM, 'assumed'
        log(f'     WARNING: no pixel size in the TIFF ({cell_roi.read_pixel_size(green_path)[1]}); assuming {px} um/px. '
            'Enter the real value (Pixel size column or Options) if it differs.')
    k = float(np.clip(cell_roi.REF_PX_UM / px, 0.2, 5))
    a = _as_rgb(a, g)
    nuc_img, _ = cell_roi.load_nuclei(blue_path, k)  # scale bar removed
    nuc, dim, ninfo = nuclei.detect(nuc_img, k, p.nuclei_method, p.thr_nuclei)
    red, _ = mina.load_channel(red_path, 0)
    green, _ = mina.load_channel(green_path, 1)
    if not (red.shape == g.shape == nuc_img.shape):
        raise ValueError(f'Image sizes differ: red {red.shape}, green {g.shape}, blue {nuc_img.shape}')
    gn, norm = cell_roi.normalize_green(g, k)
    log(f'     {red.shape[1]}x{red.shape[0]} px, pixel size {px:.4f} um ({px_source}), {nuc.max()} in-focus nuclei; '
        f'green auto-levels: background {norm["offset"]:.2f}, gain x{norm["gain"]:.2f}')

    log('2/5  Segmenting cells (one nucleus per cell; mito-free lines and cell shape set the borders)')
    n_focus = int(nuc.max())
    if p.nuclei_method != 'otsu':
        log(f"     nuclei ({p.nuclei_method}): {n_focus} in focus, {int((np.unique(dim) > 0).sum())} dim, "
            f"{ninfo['n_rejected']} smooth (untextured) blue regions rejected")
    if p.dim_nuclei:
        nuc = np.where(dim > 0, dim, nuc)
        if nuc.max() > n_focus:
            log(f'     + {int(nuc.max()) - n_focus} dim (out-of-focus) nuclei used as cell seeds')
    res = cell_roi.segment(gn, min_area_px=int(p.min_area_um2 / px ** 2), nuclei=nuc, k=k)
    res['nuc_img'] = nuc_img
    morph = cell_roi.segment_morph(red, nuc, res['landscape'], px, k, fg_level=p.cell_fg_level)
    res['labels'], res['fg'] = morph['labels'], morph['fg']
    trimmed = {}
    if p.trim_edge_cells:
        lab0 = res['labels']
        res['labels'], trimmed = cell_roi.trim_edge_cells(lab0, nuc, morph['cost'], morph['mito_density'], px)
        res['trimmed'] = (lab0 > 0) & (res['labels'] == 0)
        if trimmed:
            log(f'     {len(trimmed)} cells at the frame kept by cutting their tip off along a mito-free line')
    return dict(a=a, g=g, gn=gn, norm=norm, px=px, px_source=px_source, k=k, nuc_img=nuc_img, nuc=nuc,
                ninfo=ninfo, n_focus=n_focus, red=red, green=green, res=res, morph=morph, trimmed=trimmed)


def review_image(st):
    """RGB uint8 image for the ROI review: mito (red), green, nuclei (blue), each auto-contrasted."""
    def nz(x, p=99.5):
        x = x.astype(float)
        return np.clip(x / max(np.percentile(x, p), 1e-9), 0, 1) ** 0.7
    return (np.dstack([nz(st['red']), nz(st['gn']) * 0.8, nz(st['nuc_img'])]) * 255).astype(np.uint8)


def detect_rois(red_path, green_path, blue_path, params=None, log=print, overrides=None):
    """Cell ROIs of one image set for manual review, without the rest of the analysis.
    Returns dict(labels, nuclei, seed, image, px): `nuclei` is labelled with the cell it belongs to and `seed`
    maps every label to 'nucleus' or 'dim nucleus'."""
    p = params or Params()
    if overrides:
        p = replace(p, **{k: v for k, v in overrides.items() if k in PARAM_FIELDS})
    st = _cell_stage(red_path, green_path, blue_path, p, log)
    lab = st['res']['labels'].astype(np.int32)
    nucl = np.where(lab > 0, st['nuc'], 0).astype(np.int32)
    seed = {int(l): 'dim nucleus' if l > st['n_focus'] else 'nucleus' for l in np.unique(lab) if l > 0}
    return dict(labels=lab, nuclei=nucl, seed=seed, image=review_image(st), px=st['px'])


def _apply_roi_edit(st, edit, log):
    """Replace the detected cell ROIs by the reviewed ones (labels renumbered 1..n)."""
    old = np.asarray(edit['labels']).astype(np.int64)
    ids = [int(i) for i in np.unique(old) if i > 0]
    lut = np.zeros(max(ids + [0]) + 1, np.int32)
    for new, o in enumerate(ids, 1):
        lut[o] = new
    lab = lut[old]
    nuc = np.asarray(edit['nuclei']).astype(np.int64)
    nuc = np.where((nuc > 0) & (nuc < len(lut)), lut[np.clip(nuc, 0, len(lut) - 1)], 0)
    nuc = np.where(nuc == lab, nuc, 0).astype(np.int32)    # a nucleus belongs to its own cell only
    seed = {int(lut[int(o)]): s for o, s in edit.get('seed', {}).items() if 0 < int(o) < len(lut) and lut[int(o)]}
    st['res']['labels'] = lab
    st['res']['nuclei'] = nuc
    st['res']['trimmed'] = None
    st['nuc'] = nuc
    st['trimmed'] = {}
    st['seed'] = {int(l): seed.get(int(l), 'manual') for l in np.unique(lab) if l > 0}
    log(f'     reviewed cell ROIs used: {int(lab.max())} cells ('
        f"{sum(s == 'manual' for s in st['seed'].values())} drawn by hand)")


def run(red_path, green_path, blue_path, outdir, params=None, name=None, log=print, overrides=None, roi_edit=None):
    """Run the full pipeline and write all outputs to `outdir`. Returns a dict of results and file paths.
    `overrides`: per-image Params values (e.g. {'pixel_size_um': 0.1, 'thr_mito': 40}) on top of `params`.
    `roi_edit`: reviewed cell ROIs from detect_rois (labels, nuclei, seed), used instead of the detected ones."""
    p = params or Params()
    if overrides:
        p = replace(p, **{k: v for k, v in overrides.items() if k in PARAM_FIELDS})
    name = name or sample_name(red_path)
    os.makedirs(outdir, exist_ok=True)
    st = _cell_stage(red_path, green_path, blue_path, p, log)
    if roi_edit is not None:
        _apply_roi_edit(st, roi_edit, log)
        write_roi_edit(outdir, name, roi_edit)
    a, g, gn, norm, px, px_source, k = (st[x] for x in ('a', 'g', 'gn', 'norm', 'px', 'px_source', 'k'))
    nuc_img, nuc, ninfo, n_focus, red, green = (st[x] for x in ('nuc_img', 'nuc', 'ninfo', 'n_focus', 'red', 'green'))
    res, morph, trimmed = st['res'], st['morph'], st['trimmed']
    seed = st.get('seed')
    gstatus, ginfo = green_cells.classify(gn, res['labels'], nuc, k, p.green_pos_percent / 100, p.thr_green_bright)
    rois = cell_roi.save_outputs(a, gn, res, px, outdir, name, binuc_tau=p.binuc_tau,
                                 exclude_binuc=p.exclude_binucleate)
    n_ok = sum(r['status'] == 'ok' for r in rois)
    for r in rois:
        r['green_status'] = gstatus[int(r['roi'][4:])]
        r['seed'] = seed.get(int(r['roi'][4:]), 'manual') if seed else (
            'dim nucleus' if int(r['roi'][4:]) > n_focus else 'nucleus')
        r['edge_trimmed_um2'] = round(trimmed.get(int(r['roi'][4:]), 0) * px * px, 1)
    log(f'     {len(rois)} cell ROIs, {n_ok} pass QC (not touching the border'
        + (', not binucleate)' if p.exclude_binucleate else ')'))

    lab = res['labels']
    mask_fn = None
    manual = f'manual threshold {p.thr_mito:g}' if p.thr_mito > 0 else ''
    if p.mito_method == 'split':
        log(f"3/5  Mito objects ({manual or 'adaptive threshold'} + watershed split), then MiNA inside each cell")
        p0 = dict(mito_objects.P, sigma=mito_objects.P['sigma'] * k, manual_thr=p.thr_mito)
        sm = mito_objects.preprocess(red, px, p0)
        used = [int(r['roi'][4:]) for r in rois if r['status'] == 'ok' or p.include_edge_cells]
        q = mito_objects.scaled_params(sm, np.isin(lab, used), k, p0)

        def mask_fn(cell):
            fg, _, sep = mito_objects.segment_cell(sm, cell, px, q)
            return sep > 0, mito_objects.fragmentation_metrics(fg, sep, cell, px)
    elif p.thr_mito > 0:
        log(f'3/5  MiNA ({manual} on the preprocessed red) inside each cell')
        sm = mito_objects.preprocess(red, px, dict(mito_objects.P, sigma=mito_objects.P['sigma'] * k))

        def mask_fn(cell):
            return cell & (sm > p.thr_mito), {}
    else:
        log('3/5  MiNA (Otsu per cell) on the red channel inside each cell')
    rows, b, sk, end, junc = mina.analyze_cells(red, lab, rois, px, green, p.include_edge_cells,
                                                log=lambda s: log('     ' + s), mask_fn=mask_fn)

    log('4/5  Green on mitochondria (per cell, per mito object, green puncta)')
    mito_rows, puncta_rows, mlab, plab, thr = green_q.quantify(
        green, red, lab, b, rows, px, bg=norm['offset'], k=k, sensitivity=p.puncta_sensitivity, min_mito_area_um2=p.min_mito_area_um2,
        log=lambda s: log('     ' + s), puncta_thr=p.thr_puncta)
    for row in rows:
        c = int(row['cell'][4:])
        row['green_status'] = gstatus[c]
        row['green_bright_area_percent'] = 100 * ginfo['fraction'][c]
    cstat = {row['cell']: row['green_status'] for row in rows}
    for t in (mito_rows, puncta_rows):
        for row in t:
            row['green_status'] = cstat.get(row['cell'], '')
    n_pos = sum(s == 'positive' for s in cstat.values())
    log(f'     {n_pos} green-positive, {len(cstat) - n_pos} green-negative analysed cells '
        f'(bright green >= {ginfo["threshold"]:.1f} auto-levelled units)')
    corr = stats.correlation_table(rows, mito_rows)

    log('5/5  Writing results')
    mina.write_outputs(red, lab, rois, rows, b, sk, end, junc, outdir, name)
    files = dict(per_cell_csv=os.path.join(outdir, f'{name}_per_cell.csv'),
                 per_mito_csv=os.path.join(outdir, f'{name}_per_mito.csv'),
                 puncta_csv=os.path.join(outdir, f'{name}_green_puncta.csv'),
                 corr_csv=os.path.join(outdir, f'{name}_correlations.csv'),
                 xlsx=os.path.join(outdir, f'{name}_results.xlsx'),
                 overlay=os.path.join(outdir, f'{name}_mina_overlay.png'),
                 green_overlay=os.path.join(outdir, f'{name}_green_on_mito.png'),
                 cells=os.path.join(outdir, f'{name}_mina_cells.png'),
                 summary=os.path.join(outdir, f'{name}_summary.png'),
                 roiset=os.path.join(outdir, f'{name}_RoiSet.zip'),
                 roiset_filtered=os.path.join(outdir, f'{name}_RoiSet_filtered.zip'),
                 mito_roiset=os.path.join(outdir, f'{name}_mito_RoiSet.zip'),
                 puncta_roiset=os.path.join(outdir, f'{name}_green_puncta_RoiSet.zip'),
                 green_cells=os.path.join(outdir, f'{name}_green_cells.png'))
    cell_rows = merge_cell_rows(rois, rows)
    settings = dict(sample=name, red=red_path, green=green_path, blue=blue_path, pixel_size_um=px,
                    pixel_size_source=px_source, rois_reviewed=roi_edit is not None, green_background=norm['offset'], green_gain=norm['gain'],
                    nuclei_threshold=ninfo.get('threshold', float('nan')),
                    nuclei_threshold_auto=ninfo.get('threshold_auto', float('nan')),
                    cell_fg_level=p.cell_fg_level,
                    mito_threshold=p.thr_mito if p.thr_mito > 0 else 'auto (per cell)',
                    puncta_threshold=thr, green_bright_threshold=ginfo['threshold'],
                    **{f'param_{k_}': v for k_, v in asdict(p).items()})
    tables = [('cells', cell_rows), ('mito', mito_rows), ('green_puncta', puncta_rows), ('correlations', corr)]
    for (_, t), f in zip(tables, (files['per_cell_csv'], files['per_mito_csv'], files['puncta_csv'], files['corr_csv'])):
        write_csv(f, t)
    # green-positive and green-negative cells, measured and reported separately
    groups = {}
    for g_key, g_tag in GREEN_GROUPS:
        gc = [r for r in cell_rows if r['green_status'] == g_key]
        gm = [r for r in mito_rows if r['green_status'] == g_key]
        gp = [r for r in puncta_rows if r['green_status'] == g_key]
        gr = stats.correlation_table(gc, gm)
        groups[g_key] = dict(rows=gc, mito_rows=gm, puncta_rows=gp, correlations=gr)
        for label, t, full in (('per_cell', gc, cell_rows), ('per_mito', gm, mito_rows),
                               ('green_puncta', gp, puncta_rows), ('correlations', gr, corr)):
            write_csv(os.path.join(outdir, f'{name}_{label}_{g_tag}.csv'), t, full[0] if full else None)
    group_sheets = [(f'{s}_{g_tag}', groups[g_key][k]) for g_key, g_tag in GREEN_GROUPS
                    for s, k in (('cells', 'rows'), ('mito', 'mito_rows'), ('puncta', 'puncta_rows'), ('corr', 'correlations'))]
    write_xlsx(files['xlsx'], group_sheets + [(f'all_{n_}', t) for n_, t in tables]
               + [('settings', [settings])])
    render_green_cells(gn, res['labels'], ginfo, gstatus, {r['roi'] for r in rois if r['status'] == 'ok'
                       or p.include_edge_cells}, files['green_cells'])
    render_green(red, green, lab, mito_rows, b, mlab, plab, rows, files['green_overlay'])
    write_label_rois(mlab, 'm', files['mito_roiset'])
    write_label_rois(plab, 'p', files['puncta_roiset'])
    log(f'     Done. Results in {outdir}')
    return dict(name=name, outdir=outdir, rows=cell_rows, mito_rows=mito_rows, puncta_rows=puncta_rows,
                correlations=corr, rois=rois, groups=groups, **files)


GREEN_GROUPS = (('positive', 'green_pos'), ('negative', 'green_neg'))


PER_CELL_ROI_COLS = ('seed', 'edge_trimmed_um2', 'centroid_x', 'centroid_y', 'touches_border', 'weak_border_with')


ROI_EDIT_SUFFIX = '_roi_review.npz'


def write_roi_edit(outdir, name, edit):
    """Keep the reviewed ROIs next to the results, so a later run can start from them."""
    seed = edit.get('seed', {})
    np.savez_compressed(os.path.join(outdir, name + ROI_EDIT_SUFFIX), labels=np.asarray(edit['labels'], np.int32),
                        nuclei=np.asarray(edit['nuclei'], np.int32),
                        seed_labels=np.array(list(seed.keys()), np.int32), seed_kinds=np.array(list(seed.values())))


def read_roi_edit(outdir, name):
    """Reviewed ROIs saved by an earlier run (or None)."""
    f = os.path.join(outdir, name + ROI_EDIT_SUFFIX)
    if not os.path.exists(f):
        return None
    d = np.load(f)
    return dict(labels=d['labels'], nuclei=d['nuclei'],
                seed={int(l): str(s) for l, s in zip(d['seed_labels'], d['seed_kinds'])})


def merge_cell_rows(rois, rows):
    """One row per analysed cell: MiNA + green metrics, plus ROI geometry."""
    by = {r['roi']: r for r in rois}
    out = []
    for row in rows:
        d = {'cell': row['cell'], 'green_status': row['green_status'], **row}
        for k in PER_CELL_ROI_COLS:
            d[k] = by[row['cell']].get(k, '')
        out.append(d)
    return out


def _fmt(v):
    if isinstance(v, float):
        return f'{v:.3e}' if v and abs(v) < 1e-3 else f'{v:.4f}'  # small p values keep their digits
    return v


def write_csv(path, rows, cols=None):
    """`cols` = header to write when `rows` is empty (e.g. an image with no green-positive cell)."""
    with open(path, 'w', newline='') as f:
        if not rows and not cols:
            return
        w = csv.DictWriter(f, fieldnames=list(rows[0]) if rows else list(cols))
        w.writeheader()
        for d in rows:
            w.writerow({k: _fmt(v) for k, v in d.items()})


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        if not rows:
            continue
        cols = list(rows[0])
        ws.append(cols)
        for r in rows:
            ws.append([(None if isinstance(v, float) and v != v else (bool(v) if isinstance(v, np.bool_) else v))
                       for v in (r.get(c) for c in cols)])
        ws.freeze_panes = 'B2'
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(9, min(32, len(c) + 2))
    wb.save(path)


def write_label_rois(lab, prefix, path):
    """ImageJ RoiSet with one polygon per labelled object (mito objects / green puncta)."""
    import roifile
    from skimage import measure
    rois = []
    for r in measure.regionprops(lab):
        c = max(measure.find_contours(np.pad(r.image, 1).astype(float), 0.5), key=len) - 1
        c = c + np.array(r.bbox[:2])
        roi = roifile.ImagejRoi.frompoints(np.c_[c[:, 1], c[:, 0]])
        roi.roitype = roifile.ROI_TYPE.POLYGON; roi.name = f'{prefix}{r.label:04d}'
        rois.append(roi)
    if os.path.exists(path):
        os.remove(path)
    if rois:
        roifile.roiwrite(path, rois)


def render_green(red, green, lab, mito_rows, mito, mlab, plab, rows, path):
    """Red mito (grey) + green POI, mito outlines magenta, puncta on mito yellow, off mito cyan."""
    from matplotlib.figure import Figure
    from scipy import ndimage as ndi
    from skimage import segmentation
    rh, gh = np.percentile(red, 99.5) or 1, np.percentile(green, 99.8) or 1
    rgb = np.dstack([np.clip(red / rh, 0, 1) * 0.55, np.clip(green / gh, 0, 1), np.clip(red / rh, 0, 1) * 0.55])
    rgb[segmentation.find_boundaries(mlab, mode='outer')] = (1.0, 0.2, 0.9)
    on = np.isin(plab, np.unique(plab[(plab > 0) & mito]))
    pb = segmentation.find_boundaries(plab, mode='outer')
    rgb[pb & ndi.binary_dilation(on, iterations=1)] = (1.0, 1.0, 0.0)
    rgb[pb & ~ndi.binary_dilation(on, iterations=1)] = (0.0, 0.9, 1.0)
    analysed = np.isin(lab, [int(r['cell'][4:]) for r in rows])
    rgb[segmentation.find_boundaries(np.where(analysed, lab, 0), mode='inner')] = (1, 1, 1)
    fig = Figure(figsize=(12, 12), dpi=150); ax = fig.subplots()
    ax.imshow(rgb); ax.set_axis_off()
    for r in rows:
        cy, cx = ndi.center_of_mass(lab == int(r['cell'][4:]))
        ax.text(cx, cy, r['cell'][4:], color='white', fontsize=13, weight='bold', ha='center', va='center',
                bbox=dict(facecolor='black', alpha=0.5, pad=1.5, edgecolor='none'))
    ax.set_title('Green POI on mito. grey/purple: red mito, green: POI; magenta: mito objects; '
                 'yellow: green puncta on mito, cyan: puncta off mito; white: analysed cells', fontsize=9)
    fig.tight_layout(); fig.savefig(path)


def render_green_cells(gn, lab, info, status, analysed, path):
    """Bright green mask, cell ROIs and their green+ / green- calls."""
    from matplotlib.figure import Figure
    from scipy import ndimage as ndi
    v = np.clip(gn / max(np.percentile(gn, 99.5), 1e-9), 0, 1)
    rgb = np.dstack([v * 0.15, v, v * 0.15])
    rgb[info['mask']] = (1.0, 0.95, 0.2)
    fig = Figure(figsize=(12, 12), dpi=120); ax = fig.subplots()
    ax.imshow(rgb); ax.set_axis_off()
    for l in range(1, lab.max() + 1):
        m = lab == l
        if not m.any():
            continue
        pos = status[l] == 'positive'
        ax.contour(m, [0.5], colors='#ff5050' if pos else '#c8c8c8', linewidths=2.2 if pos else 1.0)
        cy, cx = ndi.center_of_mass(m)
        name = f'cell{l:02d}'
        ax.text(cx, cy, f"{l}{'+' if pos else '−'}"
                + ('' if name in analysed else '\n(not analysed)'),
                color='white', fontsize=(13 if pos else 10) if name in analysed else 8, weight='bold', ha='center', va='center',
                bbox=dict(facecolor='#b00000' if pos else 'black', alpha=0.6, pad=1.5, edgecolor='none'))
    ax.set_title(f'Green+ cells (red outline) / green− (grey). Yellow: bright green >= {info["threshold"]:.1f} '
                 f'auto-levelled units. Cell ROIs are not changed by green.', fontsize=9)
    fig.tight_layout(); fig.savefig(path)
